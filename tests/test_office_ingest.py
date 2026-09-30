"""Word, PowerPoint and Excel files as course sources (ingest/office.py):
uploaded like any file, indexed with citable locators, and readable in the
source viewer as the text ingestion saw."""

import io
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from src.backend.common.db import connection
from src.backend.ingest import worker
from src.backend.ingest.office import OfficeFileError, docx_text, pptx_pages, xlsx_pages

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Cell biology", level=1)
    document.add_paragraph("Mitochondria make ATP.")
    document.add_heading("Organelles", level=2)
    document.add_paragraph("Ribosomes build proteins.", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Term"
    table.cell(0, 1).text = "Meaning"
    table.cell(1, 0).text = "ATP"
    table.cell(1, 1).text = "energy currency"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation

    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = "Photosynthesis"
    first.placeholders[1].text_frame.text = "Light reactions happen in thylakoids"
    first.notes_slide.notes_text_frame.text = "Mention chlorophyll"
    deck.slides.add_slide(deck.slide_layouts[6])  # blank: keeps its number
    third = deck.slides.add_slide(deck.slide_layouts[1])
    third.shapes.title.text = "Calvin cycle"
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Grades"
    sheet.append(["Student", "Score"])
    sheet.append(["Ada", 97.0])
    sheet.append([None, None])
    sheet.append(["Grace", 88.5])
    book.create_sheet("Empty")
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_docx_headings_tables_and_lists_become_markdown() -> None:
    text = docx_text(_docx())
    assert "# Cell biology" in text and "## Organelles" in text
    assert "- Ribosomes build proteins." in text
    assert "Term | Meaning\nATP | energy currency" in text


def test_pptx_slides_keep_their_numbers_and_notes() -> None:
    pages = pptx_pages(_pptx())
    assert [label for label, _ in pages] == ["slide 1", "slide 2", "slide 3"]
    assert pages[0][1].splitlines()[0] == "Photosynthesis"
    assert "thylakoids" in pages[0][1] and "Notes: Mention chlorophyll" in pages[0][1]
    assert pages[1][1] == "" and pages[2][1] == "Calvin cycle"


def test_xlsx_sheets_become_pages_of_rows() -> None:
    pages = dict(xlsx_pages(_xlsx()))
    assert pages["sheet Grades"] == "Grades\nStudent | Score\nAda | 97\nGrace | 88.5"


def test_a_non_office_zip_is_a_readable_error() -> None:
    with pytest.raises(OfficeFileError):
        docx_text(b"definitely not a zip")


def _upload(client: TestClient, course_id: str, name: str, body: bytes, mime: str):
    return client.post(
        f"/courses/{course_id}/sources",
        data={"source_type": "notes"},
        files={"file": (name, body, mime)},
    )


@pytest.mark.parametrize(
    ("name", "body", "mime", "locator_type", "label"),
    [
        ("notes.docx", _docx, DOCX, "section", "§ Cell biology"),
        ("lecture.pptx", _pptx, PPTX, "slide", "slide 1"),
        ("grades.xlsx", _xlsx, XLSX, "sheet", "sheet Grades"),
    ],
)
def test_office_files_upload_index_and_open_in_the_viewer(
    client: TestClient, name: str, body, mime: str, locator_type: str, label: str
) -> None:
    course_id = client.post("/courses", json={"name": "Office"}).json()["course_id"]
    # The OS often reports an Office file as a generic type; the extension wins.
    uploaded = _upload(client, course_id, name, body(), "application/octet-stream")
    assert uploaded.status_code == 201, uploaded.text
    assert uploaded.json()["mime_type"] == mime
    source_id = uploaded.json()["source_id"]

    assert worker.process_batch(limit=5) == (1, 1)
    listed = client.get(f"/courses/{course_id}/sources").json()
    assert listed[0]["status"] == "indexed", listed[0]["error_message"]

    with connection() as conn:
        locator = conn.execute(
            "SELECT locator_type, label FROM locators WHERE source_id = ?"
            " ORDER BY start LIMIT 1",
            (UUID(source_id),),
        ).fetchone()
        chunk = conn.execute(
            "SELECT chunk_id, text FROM chunks WHERE source_id = ? LIMIT 1",
            (UUID(source_id),),
        ).fetchone()
    assert (locator["locator_type"], locator["label"]) == (locator_type, label)

    viewer = client.get(f"/courses/{course_id}/sources/{source_id}/content")
    assert viewer.status_code == 200
    assert viewer.headers["content-type"].startswith("text/plain")
    assert chunk["text"][:40] in viewer.text, "a cited passage is findable in the view"


def test_a_corrupt_office_file_fails_with_a_reason(client: TestClient) -> None:
    course_id = client.post("/courses", json={"name": "Office"}).json()["course_id"]
    broken = _upload(client, course_id, "broken.docx", b"not a zip", DOCX)
    assert broken.status_code == 201
    assert worker.process_batch(limit=5) == (1, 0)
    row = client.get(f"/courses/{course_id}/sources").json()[0]
    assert row["status"] == "failed"
    assert "Word file" in row["error_message"] or "Office file" in row["error_message"]
