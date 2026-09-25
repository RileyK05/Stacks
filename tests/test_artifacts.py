"""Artifacts (docs/plan-notebook.md §4.3): typed content, versions, saving
from a chat, citations, model edits as proposals, exports."""

from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import pytest
from fastapi.testclient import TestClient
from src.backend.artifacts import content as artifact_content
from src.backend.artifacts import edit as artifact_edit
from src.backend.artifacts.content import UnknownCitationError
from tests.conftest import configure_test_provider
from tests.factories import add_chunk, chunk_source_locator

LINEARITY = "linearity means preserving addition and scaling."
BASIS = "a basis is a linearly independent spanning set."


def _course(client: TestClient) -> tuple[str, list[UUID]]:
    course_id = client.post("/courses", json={"name": "Linear algebra"}).json()[
        "course_id"
    ]
    chunks = [add_chunk(UUID(course_id), text) for text in (LINEARITY, BASIS)]
    return course_id, chunks


def _create(
    client: TestClient, course_id: str, kind: str, **extra: Any
) -> dict[str, Any]:
    response = client.post(
        f"/courses/{course_id}/artifacts", json={"kind": kind, **extra}
    )
    assert response.status_code == 201, response.text
    created: dict[str, Any] = response.json()
    return created


def _save(
    client: TestClient, course_id: str, artifact: dict[str, Any], **changes: Any
) -> Any:
    body = {
        "base_version": artifact["version"],
        "title": artifact["title"],
        "content": artifact["content"],
        **changes,
    }
    return client.put(
        f"/courses/{course_id}/artifacts/{artifact['artifact_id']}", json=body
    )


# --- content ---


def test_citations_are_found_and_renumbered_everywhere() -> None:
    content = {
        "questions": [
            {"prompt": "Q [2]", "options": ["a", "b"], "answer": 0, "sources": [2, 3]},
            {"prompt": "R [3, 2]", "options": ["a", "b"], "answer": 1, "sources": [3]},
        ]
    }
    assert artifact_content.cited_numbers(content) == {2, 3}
    renumbered = artifact_content.renumber(content, {2: 1, 3: 2})
    assert renumbered["questions"][0]["prompt"] == "Q [1]"
    assert renumbered["questions"][1]["prompt"] == "R [2, 1]"
    assert renumbered["questions"][0]["sources"] == [1, 2]
    with pytest.raises(UnknownCitationError):
        artifact_content.renumber(content, {2: 1})


def test_compact_keeps_only_cited_chunks() -> None:
    a, b, c = uuid4(), uuid4(), uuid4()
    content, sources = artifact_content.compact({"markdown": "x [3] y [1]"}, [a, b, c])
    assert sources == [a, c] and content["markdown"] == "x [2] y [1]"
    with pytest.raises(UnknownCitationError):
        artifact_content.compact({"markdown": "[4]"}, [a, b, c])


def test_merge_appends_new_sources_and_reuses_known_ones() -> None:
    a, b, c = uuid4(), uuid4(), uuid4()
    content, sources = artifact_content.merge(
        {"markdown": "old [1], new [2]"}, [b, c], existing=[a, b]
    )
    assert sources == [a, b, c] and content["markdown"] == "old [2], new [3]"


def test_doc_sections_split_at_headings_and_join_back() -> None:
    text = "Intro\n\n# One\nbody\n## Two\nmore\n"
    sections = artifact_edit.doc_sections(text)
    assert sections == ["Intro\n\n", "# One\nbody\n", "## Two\nmore\n"]
    assert "".join(sections) == text


# --- create / save / versions ---


@pytest.mark.parametrize("kind", artifact_content.KINDS)
def test_every_kind_starts_blank_and_valid(client: TestClient, kind: str) -> None:
    course_id, _ = _course(client)
    created = _create(client, course_id, kind)
    assert created["version"] == 1 and created["title"].startswith("Untitled")
    listed = client.get(f"/courses/{course_id}/artifacts").json()
    assert [a["artifact_id"] for a in listed] == [created["artifact_id"]]


def test_saves_make_versions_and_stale_saves_are_refused(client: TestClient) -> None:
    course_id, _ = _course(client)
    doc = _create(client, course_id, "doc", title="Notes")
    first = _save(client, course_id, doc, content={"markdown": "# Notes\nOne"}).json()
    assert first["version"] == 2
    stale = _save(client, course_id, doc, content={"markdown": "other window"})
    assert stale.status_code == 409, (
        "a save based on version 1 can't overwrite version 2"
    )

    second = _save(
        client, course_id, first, content={"markdown": "# Notes\nTwo"}
    ).json()
    versions = client.get(
        f"/courses/{course_id}/artifacts/{doc['artifact_id']}/versions"
    ).json()
    assert [v["version"] for v in versions] == [3, 2, 1]
    restored = client.post(
        f"/courses/{course_id}/artifacts/{doc['artifact_id']}/versions/2/restore",
        json={"base_version": second["version"]},
    ).json()
    assert (
        restored["version"] == 4 and restored["content"]["markdown"] == "# Notes\nOne"
    )
    latest = client.get(
        f"/courses/{course_id}/artifacts/{doc['artifact_id']}/versions"
    ).json()[0]
    assert latest["note"] == "Restored version 2"


def test_invalid_content_and_stray_citations_are_refused(client: TestClient) -> None:
    course_id, _ = _course(client)
    sheet = _create(client, course_id, "sheet")
    ragged = _save(
        client,
        course_id,
        sheet,
        content={"columns": ["A", "B"], "rows": [["only one"]]},
    )
    assert ragged.status_code == 422
    doc = _create(client, course_id, "doc")
    uncited = _save(client, course_id, doc, content={"markdown": "claim [1]"})
    assert uncited.status_code == 422, "[1] names a source the doc doesn't have"


def test_sources_must_belong_to_the_course(client: TestClient) -> None:
    course_id, chunks = _course(client)
    other_course, other_chunks = _course(client)
    doc = _create(client, course_id, "doc")
    own = _save(
        client,
        course_id,
        doc,
        content={"markdown": "claim [1]"},
        sources=[str(chunks[0])],
        author="model",
    )
    assert own.status_code == 200
    foreign = _save(
        client,
        course_id,
        own.json(),
        content={"markdown": "claim [1]"},
        sources=[str(other_chunks[0])],
    )
    assert foreign.status_code == 422
    del other_course


# --- from a chat ---


def _quiz_reply() -> str:
    return json.dumps(
        {
            "reply": "Here is a quick check.",
            "item": {
                "type": "quiz",
                "title": "Linearity check",
                "questions": [
                    {
                        "prompt": "What does linearity preserve?",
                        "options": ["addition and scaling", "only order"],
                        "answer": 0,
                        "explanation": "See the definition [2].",
                        "sources": [2],
                    }
                ],
            },
        }
    )


def test_a_chat_item_saves_with_citations_pinned_to_its_chunks(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    course_id, _ = _course(client)
    configure_test_provider(monkeypatch, _quiz_reply())
    chat = client.post(f"/courses/{course_id}/conversations", json={}).json()
    turn = client.post(
        f"/courses/{course_id}/conversations/{chat['conversation_id']}/messages",
        json={"question": "Quiz me on linearity and a basis"},
    ).json()
    reply = turn["reply"]
    assert reply["answer"]["workspace"], "the stub reply carries a quiz"
    cited_chunk = reply["answer"]["chunk_ids"][1]

    saved = client.post(
        f"/courses/{course_id}/artifacts/from-message",
        json={"message_id": reply["message_id"], "item_index": 0},
    )
    assert saved.status_code == 201, saved.text
    artifact = saved.json()
    assert (artifact["kind"], artifact["title"]) == ("quiz", "Linearity check")
    assert artifact["sources"] == [cited_chunk], "only what the item cites, as [1]"
    question = artifact["content"]["questions"][0]
    assert (
        question["sources"] == [1]
        and question["explanation"] == "See the definition [1]."
    )
    assert artifact["origin"]["by"] == "chat"

    missing = client.post(
        f"/courses/{course_id}/artifacts/from-message",
        json={"message_id": reply["message_id"], "item_index": 3},
    )
    assert missing.status_code == 404


def test_citations_follow_source_order_and_report_removed_sources(
    client: TestClient,
) -> None:
    course_id, chunks = _course(client)
    doc = _create(client, course_id, "doc")
    _save(
        client,
        course_id,
        doc,
        content={"markdown": "a [1] b [2]"},
        sources=[str(chunks[1]), str(chunks[0])],
        author="model",
    )
    cited = client.get(
        f"/courses/{course_id}/artifacts/{doc['artifact_id']}/citations"
    ).json()
    assert [c["chunk_id"] for c in cited] == [str(chunks[1]), str(chunks[0])]
    assert cited[0]["citation"]["text"] == BASIS

    source_id, _ = chunk_source_locator(chunks[1])
    client.delete(f"/courses/{course_id}/sources/{source_id}")
    after = client.get(
        f"/courses/{course_id}/artifacts/{doc['artifact_id']}/citations"
    ).json()
    assert after[0]["citation"] is None and after[1]["citation"] is not None

    # The doc stays editable with its missing slot; only new sources are checked.
    current = client.get(f"/courses/{course_id}/artifacts/{doc['artifact_id']}").json()
    kept = _save(client, course_id, current, content={"markdown": "a [1] b [2] c"})
    assert kept.status_code == 200, kept.text
    assert kept.json()["sources"] == [str(chunks[1]), str(chunks[0])]
    gone = _save(
        client,
        course_id,
        kept.json(),
        sources=[str(chunks[1]), str(chunks[0]), str(uuid4())],
    )
    assert gone.status_code == 422


# --- model edits ---


def test_a_doc_edit_is_a_proposal_with_citations_merged(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    course_id, chunks = _course(client)
    doc = _create(client, course_id, "doc")
    saved = _save(
        client,
        course_id,
        doc,
        content={"markdown": "# Linearity\nIt preserves addition [1].\n"},
        sources=[str(chunks[0])],
        author="model",
    ).json()
    calls = configure_test_provider(
        monkeypatch,
        "# Linearity\nIt preserves addition and scaling [1]. A basis spans [2].\n",
    )
    proposal = client.post(
        f"/courses/{course_id}/artifacts/{doc['artifact_id']}/propose-edit",
        json={"request": "rewrite it to also say what a basis is"},
    )
    assert proposal.status_code == 200, proposal.text
    body = proposal.json()
    assert body["sources"][0] == str(chunks[0]), "existing citations keep their number"
    assert set(body["sources"]) == {str(chunks[0]), str(chunks[1])}
    assert "A basis spans [2]" in body["content"]["markdown"]
    prompt = str(calls[-1]["prompt"])
    assert "Requested change: rewrite it to also say what a basis is" in prompt
    assert prompt.index(LINEARITY) < prompt.index(BASIS), "cited material comes first"

    current = client.get(f"/courses/{course_id}/artifacts/{doc['artifact_id']}").json()
    assert current["version"] == saved["version"], "a proposal saves nothing"
    accepted = _save(
        client,
        course_id,
        current,
        content=body["content"],
        sources=body["sources"],
        author="model",
        note="rewrite it to also say what a basis is",
    )
    assert accepted.status_code == 200
    versions = client.get(
        f"/courses/{course_id}/artifacts/{doc['artifact_id']}/versions"
    ).json()
    assert versions[0]["author"] == "model"


def test_an_edit_citing_material_it_was_not_given_is_refused(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    course_id, _ = _course(client)
    doc = _create(client, course_id, "doc")
    configure_test_provider(monkeypatch, "Invented claim [9].")
    refused = client.post(
        f"/courses/{course_id}/artifacts/{doc['artifact_id']}/propose-edit",
        json={"request": "add linearity"},
    )
    assert refused.status_code == 422 and "wasn't given" in refused.json()["detail"]


def test_a_scoped_edit_changes_only_its_section(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    course_id, _ = _course(client)
    doc = _create(client, course_id, "doc")
    doc = _save(
        client, course_id, doc, content={"markdown": "# One\nfirst\n\n# Two\nsecond\n"}
    ).json()
    configure_test_provider(monkeypatch, "# Two\nshorter\n")
    body = client.post(
        f"/courses/{course_id}/artifacts/{doc['artifact_id']}/propose-edit",
        json={"request": "shorten", "scope": {"part": "section", "index": 1}},
    ).json()
    assert body["content"]["markdown"] == "# One\nfirst\n\n# Two\nshorter\n"


def test_a_sheet_edit_uses_structured_output(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    course_id, _ = _course(client)
    sheet = _create(client, course_id, "sheet")
    reply = {
        "columns": ["Term", "Meaning"],
        "rows": [["linearity", "addition and scaling [1]"]],
    }
    calls = configure_test_provider(monkeypatch, json.dumps(reply))
    body = client.post(
        f"/courses/{course_id}/artifacts/{sheet['artifact_id']}/propose-edit",
        json={"request": "a glossary of linearity terms"},
    ).json()
    assert body["content"]["columns"] == ["Term", "Meaning"]
    assert body["content"]["rows"][0][1] == "addition and scaling [1]"
    assert len(body["sources"]) == 1
    assert calls[-1]["task"] == "artifact_generation"


# --- export ---


def _export(client: TestClient, course_id: str, artifact_id: str, fmt: str) -> Path:
    response = client.post(
        f"/courses/{course_id}/artifacts/{artifact_id}/export", json={"format": fmt}
    )
    assert response.status_code == 200, response.text
    path = Path(response.json()["path"])
    assert path.is_file() and path.suffix == f".{fmt}"
    return path


def test_exports_open_in_office_formats_with_their_sources(client: TestClient) -> None:
    course_id, chunks = _course(client)
    doc = _create(client, course_id, "doc", title="Study guide")
    _save(
        client,
        course_id,
        doc,
        content={
            "markdown": "# Linearity\n\nIt **preserves** addition [1].\n\n"
            "- one\n- two\n\n"
            "| Term | Meaning |\n|---|---|\n| basis | spanning set |\n"
        },
        sources=[str(chunks[0])],
        author="model",
    )
    from docx import Document

    document = Document(str(_export(client, course_id, doc["artifact_id"], "docx")))
    texts = [p.text for p in document.paragraphs]
    assert "Linearity" in texts and "It preserves addition [1]." in texts
    assert any(t.startswith("[1] notes.txt") for t in texts), (
        "sources listed at the end"
    )
    assert document.tables[0].cell(1, 0).text == "basis"
    markdown = _export(client, course_id, doc["artifact_id"], "md").read_text(
        encoding="utf-8"
    )
    assert "## Sources" in markdown

    sheet = _create(
        client,
        course_id,
        "sheet",
        title="Terms",
        content={"columns": ["Term", "Meaning"], "rows": [["basis", "spanning set"]]},
    )
    from openpyxl import load_workbook

    book = load_workbook(
        io.BytesIO(
            _export(client, course_id, sheet["artifact_id"], "xlsx").read_bytes()
        )
    )
    assert [c.value for c in next(book.active.iter_rows(max_row=1))] == [
        "Term",
        "Meaning",
    ]
    csv_text = _export(client, course_id, sheet["artifact_id"], "csv").read_text(
        encoding="utf-8-sig"
    )
    assert csv_text.splitlines()[1] == "basis,spanning set"

    deck = _create(
        client,
        course_id,
        "slides",
        title="Week 1",
        content={
            "slides": [
                {"title": "Linearity", "body": "- adds\n- scales", "notes": "say it"}
            ]
        },
    )
    from pptx import Presentation

    presentation = Presentation(
        str(_export(client, course_id, deck["artifact_id"], "pptx"))
    )
    titles = [
        s.shapes.title.text for s in presentation.slides if s.shapes.title is not None
    ]
    assert titles == ["Week 1", "Linearity"]


def test_formats_are_checked_and_charts_lose_scripts(client: TestClient) -> None:
    course_id, _ = _course(client)
    sheet = _create(client, course_id, "sheet")
    wrong = client.post(
        f"/courses/{course_id}/artifacts/{sheet['artifact_id']}/export",
        json={"format": "pptx"},
    )
    assert wrong.status_code == 422
    chart = _create(
        client,
        course_id,
        "chart",
        content={
            "html": '<svg onload="steal()"><script>alert(1)</script>'
            '<a href="javascript:x">l</a></svg>'
        },
    )
    html = _export(client, course_id, chart["artifact_id"], "html").read_text(
        encoding="utf-8"
    )
    assert "<script" not in html and "onload" not in html and "javascript:" not in html


def test_deleting_the_course_removes_its_artifacts(client: TestClient) -> None:
    from src.backend.common.db import connection

    course_id, _ = _course(client)
    _create(client, course_id, "doc")
    client.delete(f"/courses/{course_id}")
    assert client.delete(f"/trash/{course_id}").status_code == 204
    with connection() as conn:
        counts = [
            conn.execute(f"SELECT COUNT(*) AS n FROM {t}").fetchone()["n"]
            for t in ("artifacts", "artifact_versions")
        ]
    assert counts == [0, 0]


def test_code_comments_are_not_section_headings() -> None:
    text = "# Intro\nx\n```python\n# a comment\nprint(1)\n```\n## Next\ny\n"
    sections = artifact_edit.doc_sections(text)
    assert [s.splitlines()[0] for s in sections] == ["# Intro", "## Next"]


def test_new_lines_lifted_from_a_passage_get_its_citation() -> None:
    from src.backend.artifacts import attribution

    material = [
        "Exams are closed book and no notes or computers are permitted.",
        "Attendance counts for twenty percent of the final course grade.",
    ]
    text = (
        "# Policies\n"
        "Attendance counts for twenty percent of the final grade.\n"
        "Exams are closed book: no notes or computers permitted.\n"
        "Bring snacks to every exam for good luck.\n"
        "My own line, already there.\n"
    )
    result = attribution.attach(text, material, before="My own line, already there.\n")
    lines = result.text.splitlines()
    assert lines[0] == "# Policies", "headings are left alone"
    assert lines[1].endswith("[2]") and lines[2].endswith("[1]")
    assert lines[3] == "Bring snacks to every exam for good luck."
    assert lines[4] == "My own line, already there.", (
        "the student's text is never cited"
    )
    assert (result.cited_now, result.uncited) == (2, 1)


def test_an_uncited_draft_is_cited_where_it_copies_the_material(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    course_id, chunks = _course(client)
    doc = _create(client, course_id, "doc")
    configure_test_provider(
        monkeypatch,
        "# Linearity\nLinearity means preserving addition and scaling operations.\n"
        "Remember to practise every single evening before class.\n",
    )
    body = client.post(
        f"/courses/{course_id}/artifacts/{doc['artifact_id']}/propose-edit",
        json={"request": "notes on linearity"},
    ).json()
    assert "scaling operations. [1]" in body["content"]["markdown"]
    assert body["sources"] == [str(chunks[0])]
    assert body["uncited_lines"] == 1


def test_a_chatty_preamble_is_not_part_of_the_doc() -> None:
    parsed = artifact_edit._parse(
        "doc", "Here is a short study guide on grading:\n\n# Grading\nA [1]\n"
    )
    assert parsed["markdown"] == "# Grading\nA [1]\n"
    kept = artifact_edit._parse("doc", "Here is the thing: it matters.\nMore.")
    assert kept["markdown"].startswith("Here is the thing"), (
        "only a lead-in ending in ':' goes"
    )


def test_an_addition_keeps_the_existing_text_and_inserts_only_the_new_part(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    course_id, chunks = _course(client)
    doc = _create(client, course_id, "doc")
    existing = "# My notes\nI wrote this myself, keep it exactly.\n"
    doc = _save(client, course_id, doc, content={"markdown": existing}).json()
    calls = configure_test_provider(
        monkeypatch, "## Basis\nA basis is a spanning set [1].\n"
    )
    body = client.post(
        f"/courses/{course_id}/artifacts/{doc['artifact_id']}/propose-edit",
        json={"request": "Add a section on what a basis is"},
    ).json()
    assert body["content"]["markdown"] == (
        "# My notes\nI wrote this myself, keep it exactly.\n\n"
        "## Basis\nA basis is a spanning set [1].\n"
    )
    assert body["sources"] == [str(chunks[1])]
    assert "Write only the new part to add" in str(calls[-1]["prompt"])


def test_pasted_material_is_removed_from_an_edit() -> None:
    from src.backend.artifacts import attribution

    material = [
        "The first exam is held in class on Wednesday September 30 and the second "
        "exam is held in class on Wednesday November 4 so plan ahead for both"
    ]
    pasted = (
        "## Exams\n"
        "Both exams are in class; plan ahead [1].\n"
        "[1] The first exam is held in class on Wednesday September 30 and\n"
        "the first exam is held in class on Wednesday September\n"
        "30 and the second exam is held in class on Wednesday\n"
        "November 4 so plan ahead for\n"
    )
    cleaned = attribution.strip_echo(pasted, material)
    assert cleaned == "## Exams\nBoth exams are in class; plan ahead [1].\n"
    quote = (
        "As the syllabus says: The first exam is held in class on "
        "Wednesday September 30.\n"
    )
    assert attribution.strip_echo(quote, material) == quote, "a quoted sentence stays"


def test_empty_docs_and_decks_are_always_drafted_as_additions() -> None:
    assert artifact_edit.is_addition("doc", "make this shorter", {"markdown": ""})
    assert artifact_edit.is_addition("doc", "Add a summary", {"markdown": "text"})
    assert not artifact_edit.is_addition(
        "doc", "Make this shorter", {"markdown": "text"}
    )
    assert not artifact_edit.is_addition(
        "doc", "Add clarity and rewrite it", {"markdown": "t"}
    )
    assert not artifact_edit.is_addition(
        "sheet", "Add a row", {"columns": [], "rows": []}
    )
