"""Text out of Word, PowerPoint and Excel files (the OOXML formats).

Students keep notes and slides in Office files; without this, the only way
in was to export them to PDF first. Each format maps onto the locator kinds
the rest of the pipeline already understands:

- `.docx` becomes Markdown-shaped text (Word's heading styles become `#`
  headings, tables become `|`-joined rows), so it gets the same section
  locators and table of contents a `.md` file does.
- `.pptx` becomes one page of text per slide (title, body, tables, speaker
  notes), cited as "slide N".
- `.xlsx` becomes one page per sheet, a row per line, cited as "sheet NAME".

Files are zip archives, so they are checked for size and member count
before a parser expands them: the stored bytes already passed the upload
ceiling, but a zip bomb is small on disk.
"""

from __future__ import annotations

import io
import zipfile
from datetime import date, datetime, time
from typing import Any

MAX_EXPANDED_BYTES = 512 * 1024 * 1024
MAX_MEMBERS = 20_000
# Spreadsheets can hold a million rows of numbers nobody will ask about;
# past this the sheet is cut and says so.
MAX_ROWS_PER_SHEET = 20_000
_TRUNCATED = "[sheet truncated: later rows were not indexed]"


class OfficeFileError(ValueError):
    """The file is not a readable Office document; the message says why."""


def _check_archive(raw: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
    except zipfile.BadZipFile as err:
        raise OfficeFileError("not a valid Office file (unreadable zip)") from err
    if len(members) > MAX_MEMBERS or sum(m.file_size for m in members) > (
        MAX_EXPANDED_BYTES
    ):
        raise OfficeFileError("this Office file expands to an unreasonable size")


def _heading_level(style_name: str) -> int:
    if style_name == "Title":
        return 1
    if style_name.startswith("Heading "):
        digits = style_name.removeprefix("Heading ")
        if digits.isdigit():
            return min(max(int(digits), 1), 6)
    return 0


def docx_text(raw: bytes) -> str:
    """The document body in reading order, headings marked with `#`."""
    _check_archive(raw)
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(io.BytesIO(raw))
    except Exception as err:  # python-docx raises many types for bad packages
        raise OfficeFileError(f"could not read this Word file: {err}") from err
    blocks: list[str] = []
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            style = item.style.name if item.style is not None else ""
            level = _heading_level(style or "")
            if level:
                blocks.append(f"{'#' * level} {text}")
            elif (style or "").startswith("List"):
                blocks.append(f"- {text}")
            else:
                blocks.append(text)
        elif isinstance(item, Table):
            rows: list[str] = []
            for row in item.rows:
                cells: list[str] = []
                seen: set[int] = set()
                for cell in row.cells:
                    if id(cell._tc) in seen:  # a merged cell repeats itself
                        continue
                    seen.add(id(cell._tc))
                    cells.append(" ".join(cell.text.split()))
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append("\n".join(rows))
    return "\n\n".join(blocks)


def _shape_lines(shape: Any) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    lines: list[str] = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            lines.extend(_shape_lines(child))
        return lines
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
        return lines
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        lines.extend(
            paragraph.strip()
            for paragraph in shape.text_frame.text.splitlines()
            if paragraph.strip()
        )
    return lines


def pptx_pages(raw: bytes) -> list[tuple[str, str]]:
    """(label, text) per slide, in deck order. A slide with no text keeps
    its place so "slide 7" stays slide 7."""
    _check_archive(raw)
    from pptx import Presentation

    try:
        deck = Presentation(io.BytesIO(raw))
    except Exception as err:
        raise OfficeFileError(f"could not read this PowerPoint file: {err}") from err
    pages: list[tuple[str, str]] = []
    for number, slide in enumerate(deck.slides, start=1):
        lines: list[str] = []
        title = slide.shapes.title
        if title is not None and title.has_text_frame and title.text_frame.text.strip():
            # A soft line break inside a title is a vertical tab.
            lines.append(" ".join(title.text_frame.text.split()))
        for shape in slide.shapes:
            if shape.shape_id == (title.shape_id if title is not None else None):
                continue
            lines.extend(_shape_lines(shape))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                lines.append(f"Notes: {notes.text.strip()}")
        pages.append((f"slide {number}", "\n".join(lines)))
    return pages


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def xlsx_pages(raw: bytes) -> list[tuple[str, str]]:
    """(label, text) per sheet: the sheet's name, then one line per
    non-empty row with cells joined by ` | `."""
    _check_archive(raw)
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as err:
        raise OfficeFileError(f"could not read this Excel file: {err}") from err
    pages: list[tuple[str, str]] = []
    try:
        for sheet in workbook.worksheets:
            lines = [sheet.title]
            for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if number > MAX_ROWS_PER_SHEET:
                    lines.append(_TRUNCATED)
                    break
                cells = [_cell(value) for value in row]
                while cells and not cells[-1]:
                    cells.pop()
                if cells:
                    lines.append(" | ".join(cells))
            pages.append((f"sheet {sheet.title}", "\n".join(lines)))
    finally:
        workbook.close()
    return pages
